# -*- coding: utf-8 -*-
"""OTIS 데시벨(dBA) 구현 변경 설명 DOCX 생성."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


def set_run_font(run, size=11, bold=False, color=None):
    run.font.name = "Malgun Gothic"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_heading_kr(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        set_run_font(run, size={1: 18, 2: 14, 3: 12}.get(level, 11), bold=True)
    return h


def add_para(doc, text, size=11, bold=False, space_after=6):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    p.paragraph_format.space_after = Pt(space_after)
    return p


def add_bullets(doc, items, size=11):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        set_run_font(run, size=size)


def add_code(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    run.font.size = Pt(9)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Pt(12)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        set_run_font(run, size=10, bold=True)
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(val)
            set_run_font(run, size=9)
    doc.add_paragraph()


def build():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Malgun Gothic"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    style.font.size = Pt(11)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("OTIS 데시벨(dBA) 측정 구현 — 변경·로직 상세 설명서")
    set_run_font(r, size=20, bold=True, color=RGBColor(0x1A, 0x3A, 0x5C))

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run(
        "작성일: 2026-09-15  ·  대상 브랜치: main  ·  "
        "범위: 마이크 소음(dBA) 캡처 → 채널 → 격자 → EVIMP1 파일"
    )
    set_run_font(r, size=10, color=RGBColor(0x55, 0x55, 0x55))

    # ------------------------------------------------------------------
    add_heading_kr(doc, "1. 무엇을 했는지 (한눈에)", 1)
    add_para(
        doc,
        "기존 main 계측 경로는 가속도·중력만 받아 3열(OTIS-VIB3: X Y Z) 파일을 "
        "만들고 있었고, 소음(dBA)은 살아 있지 않았다. 이번 작업으로 마이크에서 "
        "주변 소음을 주기적으로 계산하고, 센서 이벤트에 실어 Flutter까지 보낸 뒤, "
        "격자 환산 결과에 붙여 EVIMP1 4열(X Y Z 소음) 파일로 저장하도록 연결했다.",
    )
    add_para(doc, "핵심 결과", bold=True)
    add_bullets(
        doc,
        [
            "신규: NoiseCaptureHandler.kt (AudioRecord → RMS → dBFS → dBA)",
            "권한: AndroidManifest에 RECORD_AUDIO, 측정 시작 전 마이크 권한 요청 UI",
            "전송: SensorStreamHandler가 이벤트 Map에 noiseDba 키 추가",
            "Dart: NativeEvent / GridSample에 noiseDba, GridResampler에서 스텝 홀드",
            "저장: VibrationFileWriter 포맷을 OTIS-VIB3(3열) → EVIMP1(4열)로 전환",
            "실패 시: 권한 거부·초기화 실패해도 앱이 죽지 않고 소음만 0.0, 진동은 계속",
        ],
    )

    # ------------------------------------------------------------------
    add_heading_kr(doc, "2. 계산식과 기본값", 1)
    add_para(doc, "소음 한 점의 계산은 NoiseCaptureHandler 안에서만 한다.", bold=True)
    add_code(
        doc,
        "1) 마이크 PCM 16bit 모노, 44100 Hz\n"
        "2) 약 0.125초(≈5512 샘플)씩 읽어 RMS 계산\n"
        "   rms = sqrt( Σ (sample/32768)² / N )\n"
        "3) dBFS = 20 × log10( max(rms, 1e-5) )\n"
        "4) dBA  = dBFS + micDbfsToDbaOffset + calibrationOffsetDba\n"
        "5) 0.0 ~ 130.0 으로 clamp 후, 소수 첫째 자리로 반올림 → latestDba",
    )
    add_table(
        doc,
        ["파라미터", "기본값", "의미", "전달 경로"],
        [
            [
                "micDbfsToDbaOffset",
                "85.0",
                "마이크가 잰 dBFS를 사람이 듣는 dBA 근처로 옮기는 임시 오프셋 (OI-4)",
                "measuring_screen → sensor_channel → MainActivity.startCapture → NoiseCaptureHandler.start",
            ],
            [
                "calibrationOffsetDba",
                "0.0",
                "현장·기기별 추가 보정",
                "동일 (현재 UI는 0.0 고정 전달)",
            ],
            [
                "SAMPLE_RATE",
                "44100",
                "초당 오디오 샘플 수",
                "NoiseCaptureHandler 상수",
            ],
            [
                "프레임 길이",
                "0.125초",
                "한 번 RMS를 내는 구간",
                "NoiseCaptureHandler",
            ],
        ],
    )
    add_para(
        doc,
        "참고: 진짜 A-가중 필터(주파수별 가중)는 아직 없고, "
        "오프셋으로 dBA에 가깝게 맞추는 임시 방식이다. 값의 절대 정확도는 "
        "교정기로 맞추기 전까지 ‘상대·대략’으로 봐야 한다.",
    )

    # ------------------------------------------------------------------
    add_heading_kr(doc, "3. 측정이 돌아갈 때 전체 로직 (시간 순서)", 1)
    add_para(
        doc,
        "사용자가 측정 화면에 들어와 카운트다운이 끝나면 "
        "_initCaptureAndTimers()가 아래 순서로 돈다.",
    )

    add_heading_kr(doc, "3.1 Flutter UI — measuring_screen.dart", 2)
    add_bullets(
        doc,
        [
            "Wakelock 켜기(화면 꺼짐 방지), 경과 시간 타이머 시작",
            "requestAudioPermission() 호출 → MethodChannel "
            "'com.otis.vibration_checker/sensors_method' 의 requestAudioPermission",
            "거부 시: 빨간 SnackBar로 ‘진동만 측정, 소음 열은 0.0’ 안내 (측정은 계속)",
            "checkSensorsAvailable() → 가속도·중력 둘 다 있으면 true",
            "startCapture(calibrationOffsetDba: 0.0, micDbfsToDbaOffset: 85.0) 호출",
            "nativeEventStream.listen → 이벤트마다 GridResampler.onEvent(event)",
            "측정 종료 시 resample() → VibrationFileWriter.write / writeMeta",
        ],
    )

    add_heading_kr(doc, "3.2 Dart 채널 — sensor_channel.dart", 2)
    add_bullets(
        doc,
        [
            "requestAudioPermission: invokeMethod → bool (승인 여부)",
            "startCapture: invokeMethod('startCapture', {calibrationOffset, micDbfsToDbaOffset})",
            "EventChannel '…/sensors_stream' 으로 배치(List<Map>) 수신",
            "각 Map을 NativeEvent.fromChannelMap으로 파싱 (noiseDba 포함)",
            "깨진 Map은 droppedMapCount만 올리고 버림",
            "stopCapture: 네이티브 stopCapture → 원본 파일 경로, 스트림 캐시 비움",
        ],
    )

    add_heading_kr(doc, "3.3 Android 진입 — MainActivity.kt", 2)
    add_bullets(
        doc,
        [
            "configureFlutterEngine에서 NoiseCaptureHandler + SensorStreamHandler(소음 참조) 생성",
            "requestAudioPermission: 이미 허용이면 true, 아니면 시스템 대화상자 "
            "(REQ_AUDIO_PERMISSION=1001) 후 onRequestPermissionsResult에서 콜백 응답",
            "startCapture: 인자에서 오프셋 읽고, begin 람다에서 "
            "noiseCaptureHandler.start(...) 후 sensorStreamHandler.start()",
            "Android 12+ 고속 샘플링 권한 없으면 먼저 요청(REQ_HIGH_RATE_PERMISSION=1002)하고 "
            "승인/거부와 무관하게 begin 실행(저속이라도 측정)",
            "stopCapture: sensorStreamHandler.stop() 경로 반환 → noiseCaptureHandler.stop()",
            "onDestroy: 센서·소음 핸들러 모두 stop (자원 누수 방지)",
        ],
    )

    add_heading_kr(doc, "3.4 소음 스레드 — NoiseCaptureHandler.kt", 2)
    add_bullets(
        doc,
        [
            "start: 기존 녹음 중이면 stop 후 재시작",
            "RECORD_AUDIO 없으면 latestDba=0.0 로그만 남기고 return (예외 안 던짐)",
            "AudioRecord(MIC, 44100, MONO, PCM16) 생성·startRecording",
            "배경 스레드 OTIS_NoiseCaptureThread: 0.125초분 read → RMS → dBA → latestDba",
            "latestDba는 @Volatile — 센서 콜백 스레드가 읽어도 안전하게 최신값 공유",
            "1초에 한 번 Logcat에 dBFS/dBA/Offset 통계",
            "stop: isRecording=false, 스레드 join(300ms), AudioRecord stop/release, latestDba=0",
        ],
    )

    add_heading_kr(doc, "3.5 센서 스트림 — SensorStreamHandler.kt", 2)
    add_bullets(
        doc,
        [
            "생성자에 NoiseCaptureHandler를 주입받아 참조만 한다 (직접 마이크를 안 연다)",
            "가속도·중력을 SENSOR_DELAY_FASTEST로 전용 HandlerThread에서 수신",
            "onSensorChanged마다 Map 작성: type, tsUs, xMg, yMg, zMg, dtUs, "
            "noiseDba = noiseCaptureHandler.latestDba",
            "즉 센서 한 점이 올 때마다 ‘그때의 최신 마이크 dBA’를 도장처럼 찍는다",
            "accel/gravity 모두 같은 latestDba를 실음. 격자에서는 raw(accel) 쪽 값만 사용",
            "32개씩 모아 EventChannel로 Flutter에 배치 전송",
            "원본 raw_native_*.txt에는 기존처럼 type tsUs x y z dtUs만 기록 (noise 열 없음)",
        ],
    )

    add_heading_kr(doc, "3.6 Dart 모델 — NativeEvent", 2)
    add_bullets(
        doc,
        [
            "noiseDba: double? 추가 (채널에 키 없으면 null)",
            "fromChannelMap: noiseDba가 있으면 num→double, 타입이 이상하면 전체 null "
            "(잘못된 값 섞임 방지). 키가 없으면 null로 통과",
            "fromRecordLine / toRecordLine(원본 txt 6필드)은 그대로 — 파일 재현 경로와 분리",
        ],
    )

    add_heading_kr(doc, "3.7 격자 환산 — GridResampler / GridSample", 2)
    add_bullets(
        doc,
        [
            "측정 중: onEvent로 accel/gravity NativeEvent를 시간순 버퍼에 쌓음",
            "측정 후 resample(): 목표 격자(예: 256Hz) 시각마다 "
            "raw·gravity를 선형 보간한 뒤 (raw − gravity) = 진동 XYZ",
            "소음은 보간하지 않음. raw 쪽 valueAt이 ‘직전 실측 accel’의 noiseDba를 "
            "그대로 유지(스텝 홀드)해 GridSample.noiseDba에 넣음",
            "이유: 마이크는 ~8Hz(0.125초)로 갱신되고 센서는 수백 Hz라, "
            "소음을 XYZ처럼 선형 보간하면 가짜 중간값이 생김",
        ],
    )

    add_heading_kr(doc, "3.8 파일 저장 — VibrationFileWriter", 2)
    add_bullets(
        doc,
        [
            "formatId: 'EVIMP1' (이전 OTIS-VIB3)",
            "columnCount: 4",
            "1줄 EVIMP1, 2줄 샘플레이트(Hz), 3줄~ 각 행: X Y Z noise (공백 구분, \\r\\n)",
            "값 표기는 소수점 최대 3자리·뒤 0 제거 (기존 formatValue 동일)",
            "집계 파일(writeMeta)의 format 줄에도 EVIMP1 (4열)로 표기",
        ],
    )

    # ------------------------------------------------------------------
    add_heading_kr(doc, "4. 파일별 변경 목록 (하나하나)", 1)

    add_heading_kr(
        doc,
        "4.1 [신규] android/.../NoiseCaptureHandler.kt",
        2,
    )
    add_para(doc, "유형: 새 파일 추가", bold=True)
    add_para(
        doc,
        "역할: 마이크 AudioRecord로 짧은 구간 RMS를 내고 dBA로 바꿔 latestDba에 보관. "
        "권한/초기화 실패 시 앱 크래시 없이 0.0.",
    )
    add_para(doc, "주요 멤버", bold=True)
    add_bullets(
        doc,
        [
            "latestDba: 바깥(SensorStreamHandler)이 읽는 최신 dBA",
            "start(calibrationOffsetDba, micDbfsToDbaOffset=85.0)",
            "stop()",
            "내부 스레드명 OTIS_NoiseCaptureThread",
        ],
    )

    add_heading_kr(
        doc,
        "4.2 [수정] android/.../AndroidManifest.xml",
        2,
    )
    add_para(doc, "유형: 권한 한 줄 추가", bold=True)
    add_code(doc, '<uses-permission android:name="android.permission.RECORD_AUDIO" />')
    add_para(
        doc,
        "없으면 AudioRecord 생성 시 SecurityException 또는 권한 체크에서 "
        "소음이 항상 0.0이 된다. 스토어/설치 시 마이크 권한 고지도 이 선언이 필요.",
    )

    add_heading_kr(
        doc,
        "4.3 [수정] android/.../MainActivity.kt",
        2,
    )
    add_para(doc, "유형: 소음 핸들러 배선 + 권한 API + start/stop 연동", bold=True)
    add_bullets(
        doc,
        [
            "필드 추가: noiseCaptureHandler, permissionCallback, REQ_AUDIO_PERMISSION",
            "SensorStreamHandler(this, noiseCaptureHandler)로 생성 변경",
            "메서드 핸들러 분기 추가: requestAudioPermission",
            "startCapture: calibrationOffset / micDbfsToDbaOffset 읽고 "
            "noise.start 후 sensor.start",
            "stopCapture / onDestroy에서 noise.stop()",
            "onRequestPermissionsResult: 마이크·고속샘플링 요청코드 분기",
        ],
    )

    add_heading_kr(
        doc,
        "4.4 [수정] android/.../SensorStreamHandler.kt",
        2,
    )
    add_para(doc, "유형: 생성자 인자 추가 + 이벤트 Map 키 추가", bold=True)
    add_bullets(
        doc,
        [
            "생성자: NoiseCaptureHandler noiseCaptureHandler 추가",
            "sampleMap에 \"noiseDba\" to noiseCaptureHandler.latestDba 추가",
            "그 외 센서 등록·배치·원본 파일 기록 로직은 기존과 동일",
        ],
    )
    add_code(
        doc,
        'val sampleMap = mapOf(\n'
        '  "type" to type, "tsUs" to tsUs,\n'
        '  "xMg" to xMg, "yMg" to yMg, "zMg" to zMg, "dtUs" to dtUs,\n'
        '  "noiseDba" to noiseCaptureHandler.latestDba,\n'
        ")",
    )

    add_heading_kr(
        doc,
        "4.5 [수정] lib/adapter/sensor_channel.dart",
        2,
    )
    add_para(doc, "유형: Dart↔Native 다리 API 확장", bold=True)
    add_bullets(
        doc,
        [
            "신규 함수 requestAudioPermission() → Future<bool>",
            "startCapture({calibrationOffsetDba=0.0, micDbfsToDbaOffset=85.0}) "
            "인자 추가 후 Map으로 invokeMethod",
            "nativeEventStream은 그대로 NativeEvent로 파싱 "
            "(모델이 noiseDba를 읽게 바뀜)",
        ],
    )

    add_heading_kr(
        doc,
        "4.6 [수정] lib/ui/features/measure/measuring_screen.dart",
        2,
    )
    add_para(doc, "유형: 측정 시작 UX에 마이크 권한·오프셋 연결", bold=True)
    add_bullets(
        doc,
        [
            "_initCaptureAndTimers 안에서 센서 확인 전 requestAudioPermission()",
            "거부 시 SnackBar (AppColors.red / AppText.body)",
            "startCapture(calibrationOffsetDba: 0.0, micDbfsToDbaOffset: 85.0)",
            "이후 구독·격자 적재·종료 저장 흐름은 기존과 동일",
        ],
    )

    add_heading_kr(
        doc,
        "4.7 [수정] lib/domain/capture/native_event.dart",
        2,
    )
    add_para(doc, "유형: 이벤트 모델에 소음 필드", bold=True)
    add_bullets(
        doc,
        [
            "필드 double? noiseDba",
            "생성자·fromChannelMap 반영",
            "원본 텍스트 6필드 encode/decode는 변경 없음",
        ],
    )

    add_heading_kr(
        doc,
        "4.8 [수정] lib/domain/capture/grid_resampler.dart",
        2,
    )
    add_para(doc, "유형: 격자 행에 소음 부착 + 홀드 보간", bold=True)
    add_bullets(
        doc,
        [
            "GridSample에 noiseDba (기본 0.0)",
            "_ChannelCursor.valueAt 반환 타입에 noiseDba 추가 "
            "(직전 이벤트 before.noiseDba ?? 0.0)",
            "resample 루프에서 GridSample(..., noiseDba: rawPoint.noiseDba)",
        ],
    )

    add_heading_kr(
        doc,
        "4.9 [수정] lib/adapter/vibration_file_writer.dart",
        2,
    )
    add_para(doc, "유형: 산출 파일 포맷을 4열 EVIMP1로 전환", bold=True)
    add_table(
        doc,
        ["항목", "이전", "이후"],
        [
            ["formatId", "OTIS-VIB3", "EVIMP1"],
            ["columnCount", "3", "4"],
            ["데이터 행", "X Y Z", "X Y Z noiseDba"],
        ],
    )
    add_para(
        doc,
        "주석에도 ‘소음 열이 붙으면 EVIMP1로 바꾼다’고 이미 적혀 있었고, "
        "그 약속을 이번에 실행한 것이다. 판독 측이 EVIMP1=4열을 가정하므로 "
        "식별자와 열 수를 같이 맞춰야 한다.",
    )

    add_heading_kr(
        doc,
        "4.10 [수정] test/domain/capture/native_event_test.dart",
        2,
    )
    add_bullets(
        doc,
        [
            "기존 ‘여분 키 noiseDba 무시’ 테스트를 ‘noiseDba를 담는다’로 변경",
            "noiseDba 키가 없어도 필수 키만 있으면 파싱되고 noiseDba==null 인 테스트 추가",
        ],
    )

    # ------------------------------------------------------------------
    add_heading_kr(doc, "5. 데이터 흐름 다이어그램 (문자)", 1)
    add_code(
        doc,
        "[사용자] 측정 시작\n"
        "   │\n"
        "   ▼\n"
        "measuring_screen\n"
        "   ├─ requestAudioPermission ──► MainActivity ──► 시스템 마이크 권한\n"
        "   └─ startCapture(0.0, 85.0) ──► MainActivity\n"
        "                                    ├─ NoiseCaptureHandler.start\n"
        "                                    │     └─ 스레드: AudioRecord → latestDba\n"
        "                                    └─ SensorStreamHandler.start\n"
        "                                          └─ onSensorChanged\n"
        "                                                Map + noiseDba=latestDba\n"
        "                                                      │ EventChannel 배치\n"
        "                                                      ▼\n"
        "                                           SensorChannelManager\n"
        "                                                      │\n"
        "                                                      ▼\n"
        "                                              NativeEvent(noiseDba)\n"
        "                                                      │\n"
        "                                                      ▼\n"
        "                                              GridResampler.onEvent\n"
        "   측정 종료 ──► resample() ──► GridSample(x,y,z,noiseDba)\n"
        "                                    │\n"
        "                                    ▼\n"
        "                         VibrationFileWriter.encode\n"
        "                           EVIMP1 / rate / X Y Z dBA …",
    )

    # ------------------------------------------------------------------
    add_heading_kr(doc, "6. 권한·실패 시 동작 (안전 규칙)", 1)
    add_table(
        doc,
        ["상황", "소음", "진동 측정", "사용자에게"],
        [
            [
                "마이크 권한 허용",
                "latestDba 갱신",
                "정상",
                "권한 창만 한 번",
            ],
            [
                "마이크 권한 거부",
                "항상 0.0",
                "정상 진행",
                "SnackBar로 소음 제외 안내",
            ],
            [
                "AudioRecord 초기화 실패",
                "0.0",
                "정상",
                "Logcat 에러 (화면 안 막음)",
            ],
            [
                "고속 샘플링 권한 거부 (A12+)",
                "권한 있으면 정상",
                "저속으로라도 시작",
                "Logcat 경고",
            ],
            [
                "가속도/중력 센서 없음",
                "start 안 함",
                "시작 안 함",
                "기존 센서 실패 안내",
            ],
        ],
    )

    # ------------------------------------------------------------------
    add_heading_kr(doc, "7. 산출물·원본 파일 차이", 1)
    add_para(doc, "원본 네이티브 기록 (디버그·재현용)", bold=True)
    add_para(
        doc,
        "경로 예: 앱 외장 filesDir/raw_native_<timestamp>.txt\n"
        "형식: # 주석 + 줄마다 type tsUs xMg yMg zMg dtUs\n"
        "소음 열은 아직 안 넣음. 소음은 EventChannel 경로로만 격자·EVIMP1에 반영.",
    )
    add_para(doc, "격자 값 파일 (분석·공유용)", bold=True)
    add_para(
        doc,
        "VibrationFileWriter가 쓰는 EVIMP1 텍스트.\n"
        "예:\n"
        "EVIMP1\n"
        "256\n"
        "12.3 -4.5 1.2 45.2\n"
        "…",
    )

    # ------------------------------------------------------------------
    add_heading_kr(doc, "8. 아직 안 한 것 / 알아둘 점", 1)
    add_bullets(
        doc,
        [
            "결과 화면(result)의 실시간 소음 그래프·noiseMax 집계 엔진 연동은 "
            "이번 범위에서 깊게 안 건드림. 파일 4열까지가 1차 목표",
            "A-가중 필터·주파수 분석은 없음 (오프셋 방식)",
            "calibrationOffset을 현장 UI에서 바꾸는 화면은 아직 없음 (0.0 고정)",
            "iOS 마이크 경로는 이번 Android 중심 작업에 포함하지 않음",
            "SensorSample(model) 쪽 noiseDba 필드는 main 모델에 없을 수 있음 — "
            "현재 라이브 경로의 소음 저장은 GridSample → EVIMP1",
        ],
    )

    # ------------------------------------------------------------------
    add_heading_kr(doc, "9. 변경 파일 체크리스트", 1)
    add_table(
        doc,
        ["구분", "경로"],
        [
            ["신규", "android/app/src/main/kotlin/com/otis/vibration_checker/NoiseCaptureHandler.kt"],
            ["수정", "android/app/src/main/AndroidManifest.xml"],
            ["수정", "android/app/src/main/kotlin/com/otis/vibration_checker/MainActivity.kt"],
            ["수정", "android/app/src/main/kotlin/com/otis/vibration_checker/SensorStreamHandler.kt"],
            ["수정", "lib/adapter/sensor_channel.dart"],
            ["수정", "lib/ui/features/measure/measuring_screen.dart"],
            ["수정", "lib/domain/capture/native_event.dart"],
            ["수정", "lib/domain/capture/grid_resampler.dart"],
            ["수정", "lib/adapter/vibration_file_writer.dart"],
            ["수정", "test/domain/capture/native_event_test.dart"],
        ],
    )

    add_heading_kr(doc, "10. 실기기 확인 방법 (짧게)", 1)
    add_bullets(
        doc,
        [
            "앱 설치 후 측정 시작 → 마이크 권한 허용",
            "Logcat 필터: NoiseCaptureHandler → 1초마다 dBFS / dBA 로그",
            "측정 종료 후 생성된 EVIMP1 파일 열어 4번째 열이 0이 아닌지 확인 "
            "(조용한 방 vs 손뼉 등)",
            "권한 거부 후 재측정 → SnackBar + 4열 전부 0.0, 진동 XYZ는 정상인지 확인",
        ],
    )

    out_dirs = [
        Path(r"C:\OTIS\docs"),
        Path(r"C:\Users\박희정\Downloads"),
    ]
    name = "OTIS_데시벨_dBA_구현_변경설명서.docx"
    saved = []
    for d in out_dirs:
        d.mkdir(parents=True, exist_ok=True)
        path = d / name
        doc.save(str(path))
        saved.append(str(path))
    return saved


if __name__ == "__main__":
    paths = build()
    for p in paths:
        print("SAVED:", p)
