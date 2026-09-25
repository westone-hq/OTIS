# -*- coding: utf-8 -*-
"""EVIMP1 진동·소음 측정 쉬운 설명 DOCX 생성 (현재 구현 기준)."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


def set_run_font(run, size=11, bold=False, color=None):
    run.font.name = "Malgun Gothic"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def add_heading_kr(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        set_run_font(run, size={1: 16, 2: 13, 3: 12}.get(level, 11), bold=True)
    return h


def add_para(doc, text, size=11, bold=False, space_after=6):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    p.paragraph_format.space_after = Pt(space_after)
    return p


def add_bullets(doc, items, size=11):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        set_run_font(run, size=size)


def add_code(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    run.font.size = Pt(9)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Pt(12)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        set_run_font(run, size=10, bold=True)
    for r_i, row in enumerate(rows):
        for c_i, val in enumerate(row):
            cell = table.rows[r_i + 1].cells[c_i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(val)
            set_run_font(run, size=9)
    doc.add_paragraph()


def build():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Malgun Gothic"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Malgun Gothic")
    style.font.size = Pt(11)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("EVIMP1 진동·소음 측정 흐름 쉬운 설명")
    set_run_font(r, size=18, bold=True, color=RGBColor(0x1A, 0x3A, 0x5C))

    add_para(
        doc,
        "마이크 44.1kHz → Fast창 0.125초(슬라이딩) → 1/256초마다 RMS → "
        "dBFS → +85 → 256Hz 격자(선형보간) → EVIMP1",
        size=10,
        space_after=12,
    )

    # 1
    add_heading_kr(doc, "1. 최종적으로 무엇을 만들려는가")
    add_para(doc, "엘리베이터 안에서 휴대폰 하나로 두 가지를 동시에 측정합니다.")
    add_bullets(
        doc,
        [
            "진동: 휴대폰 가속도·중력 센서로 X, Y, Z 진동값 측정",
            "소음: 휴대폰 마이크로 주변 소리 크기를 계산",
        ],
    )
    add_para(
        doc,
        "마지막 파일에는 한 줄마다 X 진동, Y 진동, Z 진동, 소음값을 함께 넣습니다. "
        "이 4열 형식이 현재 구현에서 EVIMP1입니다.",
    )
    add_para(doc, "예시", bold=True, space_after=2)
    add_code(
        doc,
        "EVIMP1\n"
        "256\n"
        "12.300   -4.500   1.200   41.301",
    )

    # 2
    add_heading_kr(doc, "2. 마이크는 처음에 무엇을 받아오는가")
    add_para(
        doc,
        "휴대폰 마이크가 처음부터 “41 dB”처럼 소음값 하나를 주는 것은 아닙니다. "
        "실제 소리의 파형을 아주 빠르게 숫자로 바꿔서 줍니다.",
    )
    add_para(doc, "예: 132, 245, 103, -86, -231, -110 ...")
    add_para(doc, "이런 하나하나의 숫자를 PCM 샘플이라고 보면 됩니다.")
    add_para(
        doc,
        "즉 PCM 샘플은 “아주 짧은 순간의 소리 파형을 숫자로 찍은 값”입니다.",
    )

    # 3
    add_heading_kr(doc, "3. 왜 44,100Hz로 받는가")
    add_para(
        doc,
        "현재 앱은 마이크를 44,100Hz로 설정합니다. "
        "44,100Hz는 1초에 마이크 값을 44,100개 받는다는 뜻입니다.",
    )
    add_para(doc, "1초 동안 → PCM 샘플 44,100개")
    add_para(
        doc,
        "왜?\n"
        "사람이 들을 수 있는 높은 소리는 대략 20,000Hz 정도까지이고, "
        "디지털로 그 주파수를 제대로 표현하려면 최소 약 2배 빠르게 측정해야 합니다.\n"
        "20,000 × 2 ≈ 40,000Hz 이상\n"
        "그중 44,100Hz는 오디오 분야에서 널리 쓰이는 표준적인 값이라 "
        "사용하기 편하고 호환성이 좋습니다.",
    )
    add_para(
        doc,
        "중요: 44,100Hz는 엘리베이터 진동 256Hz와 맞추기 위해 고른 값이 아닙니다. "
        "44,100Hz는 마이크 원본 소리를 받는 속도이고, "
        "256Hz는 나중에 최종 진동 표를 만드는 속도입니다.",
    )
    add_para(
        doc,
        "또한 반드시 44,100Hz만 가능한 것은 아닙니다. "
        "48,000Hz 같은 다른 오디오 샘플링 속도도 사용할 수 있습니다.",
    )

    # 4
    add_heading_kr(doc, "4. 왜 0.125초 창을 쓰는가")
    add_para(
        doc,
        "1초에 44,100개의 PCM 샘플이 들어오지만, 이 44,100개가 각각 소음 dB 값인 것은 아닙니다. "
        "한 샘플은 너무 짧은 순간의 파형값이라 그것만 보고 “지금 50dB다”라고 할 수 없습니다.",
    )
    add_para(
        doc,
        "그래서 짧은 시간 동안 들어온 여러 샘플을 한 묶음(창)으로 보고, "
        "그 구간의 전체적인 소리 크기를 계산합니다.",
    )
    add_para(
        doc,
        "현재 창 길이: 0.125초 = 125ms\n"
        "44,100 × 0.125 ≈ 5,512\n"
        "즉 약 5,512개의 PCM 샘플을 “지금 소리 크기”를 계산하는 창으로 씁니다.",
    )
    add_para(
        doc,
        "왜 0.125초인가? 너무 짧으면 값이 지나치게 출렁일 수 있고, "
        "너무 길면 엘리베이터의 출발·정속·감속처럼 빠르게 변하는 소음을 늦게 따라갑니다. "
        "125ms는 소음계의 FAST 응답 시간과 같은 시간 규모라 빠른 변화 확인에 적당한 편입니다. "
        "다만 현재 앱이 전문 소음계의 FAST 필터 자체를 그대로 구현한 것은 아닙니다.",
    )

    # 5 — NEW vs old doc
    add_heading_kr(doc, "5. 창을 왜 1/256초마다 미는가 (슬라이딩)")
    add_para(
        doc,
        "예전에는 0.125초 분량을 통째로 읽고 나서야 소음값 하나를 만들었습니다. "
        "그러면 1초에 약 8번만 새 값이 나오고, 파일에서는 같은 숫자가 길게 반복되기 쉽습니다.",
    )
    add_para(
        doc,
        "지금은 참고 EVIMP1 파일처럼 행마다 소음이 조금씩 바뀌도록, "
        "0.125초 창은 유지하되 창을 통째로 끊지 않고 "
        "약 1/256초(≈3.9ms, 약 172 샘플)마다 조금씩 앞으로 밉니다.",
    )
    add_para(
        doc,
        "44,100 ÷ 256 ≈ 172\n"
        "→ 약 172개 PCM 샘플이 들어올 때마다 창을 한 칸 밀고 RMS를 다시 계산\n"
        "→ 초당 약 256번 latestDba가 갱신",
    )
    add_para(
        doc,
        "비유하면, 0.125초짜리 “창문”을 벽에 붙인 채로 "
        "아주 조금씩 옆으로 밀면서 매번 “지금 창 안에 보이는 소리 세기”를 다시 재는 것입니다. "
        "창 길이는 Fast(0.125초)이고, 미는 간격은 최종 표의 256Hz와 맞춥니다.",
    )

    # 6
    add_heading_kr(doc, "6. RMS는 무엇이고 왜 계산하는가")
    add_para(
        doc,
        "창 안에 들어 있는 샘플들로 RMS를 계산합니다. "
        "RMS는 “이 창 동안 소리가 전체적으로 얼마나 강했는지”를 하나의 숫자로 나타내는 방법입니다.",
    )
    add_para(
        doc,
        "그냥 평균을 내면 안 되는 이유는 소리 파형에 +값과 -값이 같이 있기 때문입니다.",
    )
    add_para(
        doc,
        "예: +100, -100, +100, -100\n"
        "그냥 평균 → 0\n"
        "하지만 실제로는 소리가 없는 것이 아니라 계속 크게 흔들리고 있습니다.\n"
        "RMS는 값을 제곱해서 +와 -가 서로 지워지지 않게 만든 뒤 평균을 내고 다시 제곱근을 구합니다.\n"
        "결과적으로 RMS가 크면 소리가 크고, RMS가 작으면 소리가 작다고 볼 수 있습니다.",
    )

    # 7
    add_heading_kr(doc, "7. RMS를 왜 dBFS로 바꾸는가")
    add_para(
        doc,
        "RMS는 아직 우리가 익숙하게 보는 데시벨 숫자가 아닙니다. "
        "그래서 RMS를 로그 계산하여 dBFS라는 디지털 마이크 기준의 데시벨 값으로 바꿉니다.",
    )
    add_para(
        doc,
        "dBFS는 휴대폰의 디지털 마이크가 표현할 수 있는 최대 신호를 0dBFS로 보는 기준입니다. "
        "그래서 보통 값이 음수로 나옵니다.",
    )
    add_para(
        doc,
        "예:\n"
        "-5 dBFS  → 매우 큰 신호\n"
        "-20 dBFS → 중간\n"
        "-40 dBFS → 작은 신호\n"
        "-60 dBFS → 매우 작은 신호\n"
        "0에 가까울수록 마이크로 들어온 소리가 큽니다.",
    )

    # 8
    add_heading_kr(doc, "8. 왜 +85 오프셋을 더하는가")
    add_para(
        doc,
        "문제는 우리가 일반적으로 보는 소음값은 40dB, 60dB, 80dB처럼 양수인데, "
        "dBFS는 -40, -20처럼 음수라는 점입니다.",
    )
    add_para(
        doc,
        "현재 앱은 dBFS 값에 +85를 더해서 우리가 보는 소음 레벨 숫자 범위 근처로 옮깁니다.",
    )
    add_para(
        doc,
        "예:\n"
        "-43.699 dBFS + 85 = 41.301\n"
        "→ latestDba ≈ 41.301",
    )
    add_para(
        doc,
        "왜 정확히 85인가? 현재는 이 값이 dBFS를 dBA 근처로 옮기는 임시 오프셋입니다. "
        "즉 실제 소음계와 완전히 교정된 정밀값은 아닙니다.",
    )

    # 9
    add_heading_kr(doc, "9. 현재 값이 정확한 dBA는 아닌 이유")
    add_para(
        doc,
        "전문적인 dBA는 사람 귀가 주파수별로 다르게 느끼는 특성을 반영하는 "
        "A-weighting 필터를 적용해야 합니다.",
    )
    add_para(
        doc,
        "현재 구현은 A-weighting 필터가 없고, dBFS에 +85 같은 오프셋을 더해 "
        "dBA와 비슷한 범위로 맞추는 1차 구현입니다.",
    )
    add_para(
        doc,
        "따라서 지금 값은 “정확한 절대 dBA”라기보다는 "
        "엘리베이터 운행 중 소음이 상대적으로 어떻게 변하는지 비교하는 용도로 보는 것이 안전합니다.",
    )

    # 10
    add_heading_kr(doc, "10. latestDba는 무엇인가")
    add_para(
        doc,
        "창을 한 번 밀 때마다 최신 소음값 하나가 생깁니다. "
        "앱은 그 값을 latestDba라는 변수에 보관합니다.",
    )
    add_para(
        doc,
        "약 3.9ms마다 → 41.301\n"
        "그다음 3.9ms → 41.247\n"
        "그다음 3.9ms → 41.302\n"
        "...\n"
        "즉 latestDba는 약 1/256초마다 새 값으로 바뀝니다. "
        "(창 자체는 항상 직전 약 0.125초 분량을 보고 계산합니다.)",
    )
    add_para(
        doc,
        "파일에 쓸 때는 참고 EVIMP1처럼 소수 셋째 자리까지 남깁니다.",
    )

    # 11
    add_heading_kr(doc, "11. 동시에 진동 센서는 어떻게 받는가")
    add_para(
        doc,
        "마이크가 소음을 계산하는 동안 가속도 센서와 중력 센서도 동시에 매우 빠르게 데이터를 받습니다.",
    )
    add_para(
        doc,
        "진동 쪽은 시간에 따른 X, Y, Z 변화를 세밀하게 보는 것이 목적이라 "
        "소음 레벨보다 훨씬 빠르게 값이 들어올 수 있습니다.",
    )
    add_para(
        doc,
        "앱은 나중에 가속도와 중력을 이용해 진동 X, Y, Z를 만들고, "
        "최종적으로 1초에 256줄이 되도록 시간 간격을 정리합니다.",
    )

    # 12
    add_heading_kr(doc, "12. 왜 최종적으로 256Hz 격자를 만드는가")
    add_para(
        doc,
        "휴대폰 센서 이벤트는 실제로 정확히 같은 간격으로 들어오지 않습니다. "
        "예를 들어 3.2ms, 4.1ms, 3.6ms처럼 조금씩 다를 수 있습니다.",
    )
    add_para(
        doc,
        "분석하기 쉽게 만들기 위해 측정이 끝난 뒤 1초에 정확히 256개가 되도록 "
        "일정한 시간 간격으로 다시 맞춥니다.",
    )
    add_para(
        doc,
        "256Hz = 1초에 256줄\n"
        "한 줄 간격 = 1 / 256초 ≈ 3.906ms",
    )
    add_para(
        doc,
        "이 과정에서 X, Y, Z와 소음은 모두 앞뒤 실측값을 이용해 "
        "중간값을 계산하는 선형보간을 사용합니다.",
    )

    # 13
    add_heading_kr(doc, "13. 진동값에 소음값은 어떻게 붙는가")
    add_para(
        doc,
        "센서 데이터가 들어올 때마다, 그 순간 latestDba에 저장되어 있는 "
        "가장 최근 소음값을 같이 붙입니다.",
    )
    add_para(
        doc,
        "예:\n"
        "현재 latestDba = 41.301\n"
        "진동 이벤트 1 → X Y Z + 41.301\n"
        "곧 latestDba가 41.247로 바뀌면\n"
        "진동 이벤트 2 → X Y Z + 41.247\n"
        "...",
    )
    add_para(
        doc,
        "즉 센서 한 점마다 “그 시점의 최신 소음값을 도장처럼 찍는다”고 이해하면 됩니다. "
        "소음 갱신이 256Hz에 가깝기 때문에, 찍히는 값도 자주 조금씩 달라집니다.",
    )

    # 14
    add_heading_kr(doc, "14. 왜 파일에는 소음도 매 줄 조금씩 다른가")
    add_para(
        doc,
        "최종 EVIMP1 표는 1초에 256줄이기 때문에 각 줄마다 소음 칸도 하나씩 존재합니다.",
    )
    add_para(
        doc,
        "지금은 소음 latestDba도 약 256Hz로 갱신되고, "
        "격자에서는 XYZ와 같이 앞뒤 값 사이를 선형보간합니다. "
        "그래서 참고 장비 파일처럼 행마다 소음이 미세하게 바뀌어 보입니다.",
    )
    add_para(doc, "예:", bold=True, space_after=2)
    add_code(
        doc,
        "X Y Z 41.301\n"
        "X Y Z 41.247\n"
        "X Y Z 41.302\n"
        "X Y Z 41.289\n"
        "...",
    )
    add_para(
        doc,
        "주의: 그렇다고 해서 “마이크가 초당 256번, 완전히 독립된 짧은 구간을 "
        "처음부터 다시 측정한다”는 뜻은 아닙니다. "
        "매번 보는 창은 직전 약 0.125초(겹치는 구간이 많은 슬라이딩 창)입니다. "
        "창을 조금씩 밀 때마다 결과가 조금씩 달라지는 것입니다.",
    )

    # 15
    add_heading_kr(doc, "15. 왜 소음도 선형보간하는가")
    add_para(
        doc,
        "X, Y, Z는 연속적으로 빠르게 들어오는 센서값이라 "
        "앞뒤 값 사이의 중간값을 계산해도 됩니다.",
    )
    add_para(
        doc,
        "소음도 이제는 진동 격자와 비슷한 주기로 갱신되므로, "
        "앞뒤 실측 소음값 사이를 직선으로 이어도 "
        "참고 EVIMP1처럼 부드럽게 이어진 표가 됩니다.",
    )
    add_para(
        doc,
        "예전 방식(step hold)은 새 값이 나오기 전까지 같은 숫자를 반복했습니다. "
        "지금은 그 방식을 쓰지 않고, XYZ와 같은 선형보간을 씁니다.",
    )

    # 16
    add_heading_kr(doc, "16. 마지막 EVIMP1 파일은 어떻게 생기는가")
    add_code(
        doc,
        "EVIMP1\n"
        "256\n"
        "12.300   -4.500   1.200   41.301\n"
        "12.500   -4.700   1.400   41.247\n"
        "12.100   -4.200   1.000   41.302\n"
        "...",
    )
    add_para(
        doc,
        "첫 줄 EVIMP1은 파일 형식 이름, 둘째 줄 256은 최종 데이터가 1초에 256줄이라는 뜻입니다. "
        "그다음부터 각 줄은 X 진동, Y 진동, Z 진동, 소음값 순서입니다.",
    )

    # 17
    add_heading_kr(doc, "17. 숫자 4개만 기억하면 전체가 정리된다")
    add_table(
        doc,
        ["숫자", "뜻", "왜 필요한가"],
        [
            [
                "44,100Hz",
                "마이크 원본 소리를 1초에 44,100개 받음",
                "가청대역을 충분히 담고, 일반 오디오에서 널리 쓰이는 샘플링 속도이기 때문",
            ],
            [
                "0.125초",
                "약 5,512개 샘플을 한 창으로 보고 RMS 계산",
                "순간값이 아니라 짧은 구간의 전체 소리 크기를 안정적으로 보기 위해(Fast 규모)",
            ],
            [
                "1/256초 홉",
                "창을 약 172 샘플마다 밀어 latestDba 갱신",
                "참고 EVIMP1처럼 행마다 소음이 미세하게 바뀌게 맞추기 위해",
            ],
            [
                "256Hz",
                "최종 X/Y/Z/소음 표를 1초에 256줄로 정리",
                "진동 분석용 데이터를 일정한 시간 간격으로 맞추기 위해",
            ],
        ],
    )

    # 18
    add_heading_kr(doc, "18. 전체 흐름 한 번에 보기")
    add_code(
        doc,
        "[마이크]\n"
        "1초에 44,100개의 PCM 샘플 수집\n"
        "↓\n"
        "길이 0.125초(약 5,512개) 슬라이딩 창\n"
        "↓\n"
        "약 1/256초마다 창을 밀어 RMS 재계산\n"
        "↓\n"
        "dBFS 계산 → 디지털 마이크 기준 데시벨\n"
        "↓\n"
        "+85 오프셋 → dBA와 비슷한 숫자 범위로 임시 보정\n"
        "↓\n"
        "latestDba 생성 (약 256Hz로 갱신, 소수 셋째 자리)\n"
        "\n"
        "[진동 센서]\n"
        "가속도 + 중력 센서 빠르게 수집\n"
        "↓\n"
        "측정 후 256Hz 격자로 시간 정리\n"
        "↓\n"
        "진동 X Y Z 계산 (선형보간)\n"
        "\n"
        "[합치기]\n"
        "각 센서 이벤트에 그 시점의 latestDba를 붙임\n"
        "↓\n"
        "256Hz 격자에서 소음도 XYZ처럼 선형보간\n"
        "↓\n"
        "EVIMP1 저장: X Y Z noise",
    )

    # 19
    add_heading_kr(doc, "19. 교수님께 한 번에 설명하는 문장")
    add_para(
        doc,
        "“엘리베이터에서 진동과 객실 소음을 동시에 측정합니다. "
        "마이크는 44.1kHz로 원본 PCM 데이터를 받고, "
        "약 0.125초 길이의 Fast 규모 창으로 RMS를 계산합니다. "
        "이 창을 약 1/256초마다 밀어 가며 소음 레벨을 자주 갱신하고, "
        "dBFS로 바꾼 뒤 현재는 +85의 임시 오프셋을 적용해 dBA에 가까운 값으로 사용합니다. "
        "진동 데이터는 최종적으로 256Hz 격자로 맞추며, "
        "소음도 XYZ와 같이 앞뒤 값 사이를 선형보간해 "
        "참고 EVIMP1처럼 행마다 미세하게 변하는 4열(X, Y, Z, 소음) 파일을 만듭니다. "
        "현재는 A-weighting과 실제 소음계 교정이 적용되지 않아 "
        "절대적인 dBA라기보다는 1차적인 소음 레벨 근사값입니다.”",
    )

    out_docs = Path(__file__).resolve().parent / "EVIMP1_진동_소음_측정_쉬운설명.docx"
    out_dl = Path(r"c:\Users\박희정\Downloads") / "EVIMP1_진동_소음_측정_쉬운설명_현재구현.docx"
    doc.save(out_docs)
    doc.save(out_dl)
    print(f"saved: {out_docs}")
    print(f"saved: {out_dl}")


if __name__ == "__main__":
    build()
