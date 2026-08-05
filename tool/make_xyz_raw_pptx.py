"""XYZ · Raw / Gravity / Linear / Motion — 회의용 PPT v2 (폰 vs 우리 공식 구분)"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

OUT = Path("docs") / "회의자료_XYZ_raw값_코드_v2.pptx"

NAVY = RGBColor(0x1F, 0x38, 0x64)
GRAY = RGBColor(0x33, 0x33, 0x33)
SUB = RGBColor(0x55, 0x55, 0x55)
CODE_BG = RGBColor(0xF4, 0xF6, 0xF8)


def set_run(run, size=18, bold=False, color=GRAY):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "맑은 고딕"


def add_title_slide(prs, title, subtitle):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    box = slide.shapes.add_textbox(Inches(0.7), Inches(2.0), Inches(8.6), Inches(1.2))
    p = box.text_frame.paragraphs[0]
    p.text = title
    set_run(p.runs[0], 32, True, NAVY)

    box2 = slide.shapes.add_textbox(Inches(0.7), Inches(3.2), Inches(8.6), Inches(2.2))
    tf2 = box2.text_frame
    tf2.word_wrap = True
    for i, line in enumerate(subtitle.split("\n")):
        para = tf2.paragraphs[0] if i == 0 else tf2.add_paragraph()
        para.text = line if line else " "
        if para.runs:
            set_run(para.runs[0], 16, False, SUB)
    return slide


def add_bullets_slide(prs, title, bullets, footer=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tbox = slide.shapes.add_textbox(Inches(0.6), Inches(0.35), Inches(8.8), Inches(0.7))
    tp = tbox.text_frame.paragraphs[0]
    tp.text = title
    set_run(tp.runs[0], 26, True, NAVY)

    body = slide.shapes.add_textbox(Inches(0.7), Inches(1.15), Inches(8.6), Inches(5.6))
    tf = body.text_frame
    tf.word_wrap = True
    for i, b in enumerate(bullets):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.text = f"• {b}"
        para.space_after = Pt(6)
        set_run(para.runs[0], 16, False, GRAY)

    if footer:
        fbox = slide.shapes.add_textbox(Inches(0.7), Inches(6.7), Inches(8.6), Inches(0.4))
        fp = fbox.text_frame.paragraphs[0]
        fp.text = footer
        set_run(fp.runs[0], 12, False, SUB)
    return slide


def add_code_slide(prs, title, why, code_lines, footer=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    tbox = slide.shapes.add_textbox(Inches(0.6), Inches(0.28), Inches(8.8), Inches(0.55))
    tp = tbox.text_frame.paragraphs[0]
    tp.text = title
    set_run(tp.runs[0], 22, True, NAVY)

    wbox = slide.shapes.add_textbox(Inches(0.6), Inches(0.85), Inches(8.8), Inches(0.7))
    wp = wbox.text_frame.paragraphs[0]
    wp.text = why
    wbox.text_frame.word_wrap = True
    if wp.runs:
        set_run(wp.runs[0], 13, False, SUB)

    shape = slide.shapes.add_shape(1, Inches(0.55), Inches(1.5), Inches(8.9), Inches(5.2))
    shape.fill.solid()
    shape.fill.fore_color.rgb = CODE_BG
    shape.line.color.rgb = RGBColor(0xD0, 0xD0, 0xD0)

    cbox = slide.shapes.add_textbox(Inches(0.72), Inches(1.7), Inches(8.55), Inches(4.6))
    tf = cbox.text_frame
    tf.word_wrap = True
    for i, line in enumerate(code_lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.text = line if line else " "
        para.space_after = Pt(1)
        if para.runs:
            set_run(para.runs[0], 11, False, GRAY)
            para.runs[0].font.name = "Consolas"

    if footer:
        fbox = slide.shapes.add_textbox(Inches(0.7), Inches(6.65), Inches(8.6), Inches(0.4))
        fp = fbox.text_frame.paragraphs[0]
        fp.text = footer
        set_run(fp.runs[0], 11, False, SUB)
    return slide


def main():
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    add_title_slide(
        prs,
        "XYZ 센서값 — 폰 vs 우리 공식",
        "Galaxy S22 · 256Hz\n"
        "Raw / Gravity / Linear = Android\n"
        "Motion = 우리 코드 (Raw − Gravity)",
    )

    add_code_slide(
        prs,
        "1. Raw XYZ — SensorStreamHandler.kt",
        "① 버퍼 저장 → ② 256Hz 시각에 꺼내기 → ③ mg 변환 (3단계)",
        [
            "// [1] 센서 이벤트 → rawSample 버퍼 저장",
            "Sensor.TYPE_ACCELEROMETER -> {",
            "  rawSample.push(event.timestamp,",
            "    event.values[0],  // X → 버퍼",
            "    event.values[1],  // Y",
            "    event.values[2])  // Z",
            "  return",
            "}",
            "",
            "// [2] 256Hz 시각 — Linear while 루프 안",
            "val (rawX, rawY, rawZ) =",
            "    rawSample.interpolateAt(nextTargetNs)",
            "// event.values[0]이 버퍼에 있던 값을",
            "// nextTargetNs 시각으로 보간해 rawX 변수 생성",
            "",
            "// [3] m/s² → mg, Flutter Map",
            "\"rawX\" to (rawX * MPS2_TO_MG),",
            "\"rawY\" to (rawY * MPS2_TO_MG),",
            "\"rawZ\" to (rawZ * MPS2_TO_MG),",
        ],
        footer="android/.../SensorStreamHandler.kt · AxisSample.push / interpolateAt",
    )

    add_code_slide(
        prs,
        "2. Gravity XYZ — SensorStreamHandler.kt",
        "① 버퍼 저장 → ② 256Hz 시각에 꺼내기 → ③ mg 변환 (Raw와 동일 구조)",
        [
            "// [1] 센서 이벤트 → gravitySample 버퍼 저장",
            "Sensor.TYPE_GRAVITY -> {",
            "  gravitySample.push(event.timestamp,",
            "    event.values[0],  // X → 버퍼",
            "    event.values[1],  // Y",
            "    event.values[2])  // Z",
            "  return",
            "}",
            "",
            "// [2] 256Hz 시각 — 같은 while 루프 안",
            "val (gravX, gravY, gravZ) =",
            "    gravitySample.interpolateAt(nextTargetNs)",
            "",
            "// [3] m/s² → mg, Flutter Map",
            "\"gravityX\" to (gravX * MPS2_TO_MG),",
            "\"gravityY\" to (gravY * MPS2_TO_MG),",
            "\"gravityZ\" to (gravZ * MPS2_TO_MG),",
            "// Z 정지 시 gravityZ ≈ 1000 mg",
        ],
        footer="android/.../SensorStreamHandler.kt · 우리가 정하는 숫자 아님",
    )

    add_code_slide(
        prs,
        "3. Linear XYZ — SensorStreamHandler.kt",
        "① 버퍼 저장 → ② 256Hz 보간 → ③ mg 변환 (Raw는 interpolateAt, Linear는 인라인)",
        [
            "// [1] Linear 이벤트 → linearSample 버퍼 저장",
            "linearSample.push(currNs, currX, currY, currZ)",
            "// currX/Y/Z = event.values[0/1/2]",
            "",
            "// [2] 256Hz 시각 — while (nextTargetNs <= currNs)",
            "val alpha = (nextTargetNs - prevTs) / (currTs - prevTs)",
            "val interpX = linearSample.prevX",
            "    + alpha * (linearSample.currX - linearSample.prevX)",
            "// interpY, interpZ 동일",
            "",
            "// [3] m/s² → mg, Flutter Map",
            "\"x\" to (interpX * MPS2_TO_MG),  // linearX",
            "\"y\" to (interpY * MPS2_TO_MG),",
            "\"z\" to (interpZ * MPS2_TO_MG),",
            "// Android 내부 필터·퓨전 (공개 안 됨)",
        ],
        footer="android/.../SensorStreamHandler.kt · lib/.../sensor_sample.dart x,y,z",
    )

    add_code_slide(
        prs,
        "4. Motion XYZ — sensor_sample.dart (우리 공식)",
        "Kotlin Map으로 받은 rawX/gravityX(mg) → Dart에서 Motion 계산",
        [
            "// 입력: SensorSample.fromMap()으로 이미 mg 단위",
            "//  rawX, gravityX ← Kotlin [3]단계에서 전달됨",
            "",
            "// ── X축 ──",
            "motionX = rawX - gravityX",
            "  (없으면 linear x 사용)",
            "",
            "// ── Y/Z축 동일 ──",
            "motionY = rawY - gravityY",
            "motionZ = rawZ - gravityZ",
            "",
            "double get motionX =>",
            "  rawX!=null && gravityX!=null",
            "    ? rawX! - gravityX! : x;",
            "",
            "// 예: 16.27 - 14.88 = 1.39 mg",
            "// Linear 평균 = 0.79 mg (폰, 별도 채널)",
        ],
        footer="lib/domain/measure/sensor_sample.dart",
    )

    add_code_slide(
        prs,
        "5. Motion 계산 — 한 샘플 예시 (X축)",
        "코드가 하는 일을 숫자로 풀어 쓴 것",
        [
            "입력 (폰이 준 값):",
            "  rawX     = 16.27 mg",
            "  gravityX = 14.88 mg",
            "",
            "우리 코드:",
            "  motionX = rawX - gravityX",
            "          = 16.27 - 14.88",
            "          =  1.39 mg",
            "",
            "비교 (같은 구간):",
            "  linearX 평균 = 0.79 mg  ← 폰이 Linear로 준 값",
            "  motionX 평균 = 1.39 mg  ← 우리가 Raw−Gravity로 계산",
            "",
            "→ 둘 다 “중력 뺀 가속도” 목적, 숫자는 다를 수 있음",
        ],
        footer="Y축·Z축도 동일: motionY = rawY−gravityY, motionZ = rawZ−gravityZ",
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(f"Wrote {OUT.resolve()}")


if __name__ == "__main__":
    main()
